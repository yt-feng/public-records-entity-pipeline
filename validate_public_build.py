"""Fast post-build invariants for the public workbook."""

from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).parent
WORKBOOK = ROOT / "inputs/entity_network_master.xlsx"
GCC = {"Saudi Arabia", "Qatar", "United Arab Emirates", "Kuwait", "Bahrain", "Oman"}


def ids(ws, column: int) -> tuple[int, int]:
    values = [str(row[column - 1] or "").strip() for row in ws.iter_rows(min_row=2, values_only=True)]
    values = [value for value in values if value]
    return len(values), len(set(values))


def main() -> None:
    with zipfile.ZipFile(WORKBOOK) as archive:
        if "[Content_Types].xml" not in archive.namelist():
            raise SystemExit("invalid xlsx zip")

    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    required = {"People", "Relationships", "Unique_People", "GCC_Unique_People", "Sources"}
    missing = required - set(wb.sheetnames)
    if missing:
        raise SystemExit(f"missing sheets: {sorted(missing)}")

    people_count, people_unique = ids(wb["People"], 1)
    relationship_count, relationship_unique = ids(wb["Relationships"], 1)
    unique_count, unique_unique = ids(wb["Unique_People"], 1)
    gcc_ws = wb["GCC_Unique_People"]
    gcc_count, gcc_unique = ids(gcc_ws, 1)
    if people_count != people_unique or relationship_count != relationship_unique:
        raise SystemExit("duplicate observation IDs")
    if unique_count != unique_unique or gcc_count != gcc_unique:
        raise SystemExit("duplicate entity IDs")

    bad_scope = []
    for row in gcc_ws.iter_rows(min_row=2, values_only=True):
        countries = {part.strip() for part in str(row[3] or "").split("；") if part.strip()}
        if not countries <= GCC:
            bad_scope.append((row[0], sorted(countries - GCC)))
    if bad_scope:
        raise SystemExit(f"strict-scope violations: {len(bad_scope)}")

    forbidden = re.compile(r"(?:/Users/|cloud-storage|public-researcher|workstream-0[12])", re.I)
    suspicious = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                if forbidden.search(str(value or "")):
                    suspicious.append((sheet.title, str(value)[:100]))
                    break
    if suspicious:
        raise SystemExit(f"local-label leakage: {len(suspicious)}")

    print(json.dumps({
        "people_observations": people_count,
        "relationship_observations": relationship_count,
        "unique_entities": unique_count,
        "strict_scope_entities": gcc_count,
        "xlsx": str(WORKBOOK),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
