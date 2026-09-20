"""Append de-duplicated fifth-wave GCC Wikidata relationship claims."""

from __future__ import annotations

import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

from relation_state import load_relation_state


ROOT = Path(__file__).parent
WORKBOOK = ROOT / "inputs/entity_network_master.xlsx"
SOURCE_ID = "SRC-068"
GCC = {"Saudi Arabia", "Qatar", "United Arab Emirates", "Kuwait", "Bahrain", "Oman"}


def style(ws, source: int, target: int) -> None:
    for col in range(1, ws.max_column + 1):
        before, after = ws.cell(source, col), ws.cell(target, col)
        if before.has_style:
            after._style = copy(before._style)
        if before.alignment:
            after.alignment = copy(before.alignment)
        if before.protection:
            after.protection = copy(before.protection)


def main() -> None:
    records = load_relation_state().get("records", [])
    wb = load_workbook(WORKBOOK)
    people, relationships = wb["People"], wb["Relationships"]
    sources, attachment = wb["Sources"], wb["Attachment_Readout"]

    people_ids = {
        str(row[0]).strip()
        for row in people.iter_rows(min_row=2, max_row=people.max_row, max_col=19, values_only=True)
        if row[0]
    }
    adds = [
        record
        for record in records
        if record.get("record_id") not in people_ids and record.get("country_section") in GCC
    ]

    people_row, relationship_row = people.max_row, relationships.max_row
    for record in adds:
        people.append([
            record["record_id"],
            record["country_section"],
            record["house"],
            f'{record["country_section"]}/Wikidata entity claims wave 5/{record["edge"]}',
            record["name"],
            "",
            "Public structured family-network candidate; fifth-wave endpoint human check passed; current identity, role and living status not assumed",
            record["intro"],
            record["relationship"],
            record["relation_type"],
            record["record_status"],
            "Wikidata GCC entity-claim wave-5 observation",
            record["evidence_level"],
            SOURCE_ID,
            record["page_url"],
            record["as_of"],
            "Confirm both endpoint identities, edge direction, family affiliation and current status with an independent source",
            f'Source member QID={record["member_qid"]}; related endpoint QID={record["related_qid"]}; fifth-wave expansion',
            "",
        ])
        people_row += 1
        style(people, people_row - 1, people_row)

        relationships.append([
            f'WDC5-REL-{record["record_id"]}',
            f'UNRESOLVED-MEMBER:{record["member_qid"]}',
            record["member_name"],
            record["record_id"],
            record["related_name"],
            record["relation_type"],
            record["relationship"],
            "补充 GCC Wikidata 第五波显式父母、兄弟姐妹、配偶和子女 claims；不等于独立血缘证明或现任王室身份。",
            "P3",
            record["evidence_level"],
            SOURCE_ID,
            record["page_url"],
            "人工核验两个端点、关系方向、家族归属和跨来源重复",
        ])
        relationship_row += 1
        style(relationships, relationship_row - 1, relationship_row)

    people.tables["PeopleTable"].ref = f"A1:S{people.max_row}"
    relationships.tables["RelationshipTable"].ref = f"A1:M{relationships.max_row}"

    source_ids = {row[0] for row in sources.iter_rows(min_row=2, values_only=True) if row[0]}
    if SOURCE_ID not in source_ids:
        sources.append([
            SOURCE_ID,
            "开放结构化数据",
            "Wikidata GCC entity claims wave 5",
            "https://www.wikidata.org/w/api.php",
            "从当前严格 GCC QID 集合继续读取 P22/P25/P3373/P26/P40 claims，排除既有关系键",
            "低/结构化关系扩展",
            "2026-09-19",
            "第五波关系边仍需独立核验；端点 P31=human 不等于现任王室身份",
        ])
        style(sources, sources.max_row - 1, sources.max_row)
    sources.tables["SourcesTable"].ref = f"A1:H{sources.max_row}"

    attachment_ids = {row[0] for row in attachment.iter_rows(min_row=2, values_only=True) if row[0]}
    if SOURCE_ID not in attachment_ids:
        attachment.append([
            SOURCE_ID,
            "Wikidata GCC entity claims 第五波。",
            "GCC 结构化关系扩展层",
            "从当前严格 GCC QID 集合继续读取显式父母、兄弟姐妹、配偶和子女 claims，按 member/related/edge/context 排除重复",
            "否",
            "结构化 claims 是关系线索；需独立核验",
        ])
        style(attachment, attachment.max_row - 1, attachment.max_row)
    attachment.tables["AttachmentReadoutTable"].ref = f"A1:F{attachment.max_row}"

    counts = {}
    for row in people.iter_rows(min_row=2, max_row=people.max_row, max_col=19, values_only=True):
        counts[row[1]] = counts.get(row[1], 0) + 1

    dashboard = wb["Dashboard"]
    target = int(dashboard["B7"].value or 100000)
    dashboard["B4"], dashboard["B5"] = people.max_row - 1, relationships.max_row - 1
    dashboard["B8"] = max(target - dashboard["B4"].value, 0)
    dashboard["B9"] = sum(counts.get(country, 0) for country in GCC)
    dashboard["B10"] = dashboard["B4"].value - dashboard["B9"].value
    dashboard["B29"] = dashboard["B9"].value
    dashboard["B30"] = round(dashboard["B9"].value / dashboard["B27"].value, 4)
    dashboard["B31"] = round(dashboard["B9"].value / dashboard["B28"].value, 4)

    coverage = wb["Coverage"]
    for row_num in range(2, coverage.max_row + 1):
        coverage.cell(row_num, 5).value = counts.get(coverage.cell(row_num, 1).value, 0)
    coverage["J4"], coverage["J5"], coverage["J6"] = sources.max_row - 1, people.max_row - 1, relationships.max_row - 1

    wb["ReadMe"]["B5"] = (
        f"People 表目前 {people.max_row - 1:,} 条命名/候选/关系观察记录；GCC 当前为 {dashboard['B9'].value:,} 条。"
        f"本轮追加 Wikidata GCC 第五波 {len(adds):,} 条显式关系观察，并排除既有关系键；结构化关系不等于独立在世王室成员。"
        "GCC 广义家族规模采用约 45,500–80,000 的 top-down 工作区间；100,000 是记录目标，不等于 100,000 名独立在世王室成员。"
    )
    wb.save(WORKBOOK)
    print(json.dumps({
        "new_people": len(adds),
        "new_relationships": len(adds),
        "people": people.max_row - 1,
        "relationships": relationships.max_row - 1,
        "gcc_records": dashboard["B9"].value,
        "gap_to_100k": dashboard["B8"].value,
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
