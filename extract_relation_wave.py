"""Fifth-wave continuation over the currently materialized strict-GCC QID set."""

from __future__ import annotations

import hashlib
import json
import os
import re
import time
from pathlib import Path

from openpyxl import load_workbook

import extract_relation_wave_base as base


ROOT = Path(__file__).parent
WORKBOOK = ROOT / "inputs/entity_network_master.xlsx"
OUTPUT = ROOT / "data/relation_wave_records.json"
CHECKPOINT = ROOT / "data/relation_wave_checkpoint.json"
FAST_INDEX = ROOT / "data/fast_entity_index.json"
GCC = {"Saudi Arabia", "Qatar", "United Arab Emirates", "Kuwait", "Bahrain", "Oman"}
SOURCE_ID = "SRC-068"


def stable(key: tuple[str, str, str, str, str]) -> str:
    return "WDC5-" + hashlib.sha1("|".join(key).encode("utf-8")).hexdigest()[:16].upper()


def load_records(path: Path) -> list[dict]:
    return json.loads(path.read_text(encoding="utf-8"))["records"] if path.exists() else []


def load_contexts() -> dict[str, set[tuple[str, str]]]:
    """Read source contexts from the compact fast index when available.

    Fast mode intentionally leaves the reader-facing workbook unchanged between
    waves.  The index therefore becomes the source-QID frontier after its first
    initialization; the workbook remains the fallback for legacy/manual runs.
    """
    if FAST_INDEX.exists():
        state = json.loads(FAST_INDEX.read_text(encoding="utf-8"))
        contexts: dict[str, set[tuple[str, str]]] = {}
        entities = state.get("entities", {})
        for key, entity_id in state.get("key_to_entity", {}).items():
            if not key.startswith("QID:"):
                continue
            qid = key[4:]
            entity = entities.get(entity_id, {})
            pairs = entity.get("contexts", [])
            if not pairs:
                pairs = [[country, house] for country in entity.get("countries", []) for house in (entity.get("houses", []) or [""])]
            for pair in pairs:
                if not pair:
                    continue
                country = str(pair[0])
                if country not in GCC:
                    continue
                contexts.setdefault(qid, set()).add((country, str(pair[1] if len(pair) > 1 else "")))
        if contexts:
            return contexts

    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    contexts = {}
    for row in wb["People"].iter_rows(min_row=2, values_only=True):
        if row[1] not in GCC:
            continue
        match = re.search(r"wikidata\.org/(?:wiki/|entity/)(Q\d+)", str(row[14] or ""))
        if match:
            contexts.setdefault(match.group(1), set()).add((row[1], row[2]))
    wb.close()
    return contexts


def main() -> None:
    contexts = load_contexts()
    all_source_qids = sorted(contexts)
    source_offset = max(int(os.getenv("RELATION_SOURCE_OFFSET", "0")), 0)
    source_limit = max(int(os.getenv("RELATION_SOURCE_LIMIT", "300")), 0)
    source_qids = all_source_qids[source_offset:source_offset + source_limit if source_limit else None]

    previous: list[dict] = []
    for path in sorted(ROOT.glob("data/network_edge*_records.json")):
        previous.extend(load_records(path))
    previous.extend(load_records(OUTPUT))
    previous_keys = {
        (record["member_qid"], record["related_qid"], record["edge"], record["country_section"], record["house"])
        for record in previous
    }

    if CHECKPOINT.exists():
        state = json.loads(CHECKPOINT.read_text(encoding="utf-8"))
        entities = state.get("entities", {})
        done = set(state.get("done", []))
    else:
        entities, done = {}, set()
    for batch in base.batches([qid for qid in source_qids if qid not in done]):
        entities.update(base.api(batch).get("entities", {}))
        done.update(batch)
        CHECKPOINT.write_text(
            json.dumps({"entities": entities, "done": sorted(done)}, ensure_ascii=False), encoding="utf-8"
        )
        time.sleep(0.2)

    target_qids = sorted({qid for entity in entities.values() for qid, _ in base.targets(entity)})
    target_entities: dict[str, dict] = {}
    target_checkpoint = ROOT / "data/relation_wave_target_checkpoint.json"
    if target_checkpoint.exists():
        state = json.loads(target_checkpoint.read_text(encoding="utf-8"))
        target_entities = state.get("entities", {})
        target_done = set(state.get("done", []))
    else:
        target_done = set()
    for batch in base.batches([qid for qid in target_qids if qid not in target_done]):
        target_entities.update(base.api(batch).get("entities", {}))
        target_done.update(batch)
        target_checkpoint.write_text(
            json.dumps({"entities": target_entities, "done": sorted(target_done)}, ensure_ascii=False), encoding="utf-8"
        )
        time.sleep(0.2)

    records = []
    seen = set()
    for member_qid, entity in entities.items():
        for related_qid, edge in base.targets(entity):
            target = target_entities.get(related_qid, {})
            p31 = {
                claim.get("mainsnak", {}).get("datavalue", {}).get("value", {}).get("id")
                for claim in target.get("claims", {}).get("P31", [])
            }
            if "Q5" not in p31:
                continue
            for country, house in sorted(contexts.get(member_qid, set())):
                key = (member_qid, related_qid, edge, country, house)
                if key in previous_keys or key in seen:
                    continue
                seen.add(key)
                member_name, related_name = base.label(entity), base.label(target)
                records.append(
                    {
                        "record_id": stable(key),
                        "country_section": country,
                        "house": house,
                        "member_qid": member_qid,
                        "related_qid": related_qid,
                        "member_name": member_name,
                        "related_name": related_name,
                        "edge": edge,
                        "page_url": f"https://www.wikidata.org/wiki/{related_qid}",
                        "name": related_name,
                        "source_text_short": f"Wikidata wave-5 {edge} claim from GCC entity {member_name}; endpoint is a Wikidata human entity.",
                        "record_status": "public GCC family-network candidate; fifth-wave explicit relationship claim; current status not assumed",
                        "evidence_level": "低/结构化关系扩展",
                        "intro": f"Wikidata wave 5 records {member_name} -> {related_name} as {edge}; endpoint P31 includes human; independent verification required.",
                        "relationship": f"{member_name} -> {related_name}: {edge}; source context={house}; fifth-wave QID expansion.",
                        "relation_type": f"GCC家族网络第五波/{edge}",
                        "source_id": SOURCE_ID,
                        "as_of": "2026-09-19",
                    }
                )

    # Keep this file append-only across GitHub Actions runs.  The extractor emits
    # only keys not present in earlier waves, so replacing the file here would
    # discard records produced by a previous runner invocation.
    existing_records = load_records(OUTPUT)
    merged_records = []
    merged_ids = set()
    for record in existing_records + records:
        if record.get("record_id") in merged_ids:
            continue
        merged_ids.add(record.get("record_id"))
        merged_records.append(record)
    OUTPUT.write_text(
        json.dumps({"records": merged_records, "source_qids": source_qids, "all_source_qids": all_source_qids, "source_offset": source_offset, "source_limit": source_limit, "target_qids": target_qids}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    for path in [CHECKPOINT, target_checkpoint]:
        if path.exists():
            path.unlink()
    print(json.dumps({"source_qids": len(source_qids), "target_qids": len(target_qids), "new_records": len(records), "retained_records": len(merged_records), "output": str(OUTPUT)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
