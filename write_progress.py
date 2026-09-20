"""Write a small machine-readable progress marker for the public workflow."""

from __future__ import annotations

import json
import os
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from relation_state import load_relation_state


ROOT = Path(__file__).parent
WORKBOOK = ROOT / "inputs/entity_network_master.xlsx"
OUTPUT = ROOT / "data/pipeline_progress.json"


def rows(workbook, sheet_name: str) -> int:
    return max(workbook[sheet_name].max_row - 1, 0)


def main() -> None:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    relation_state = load_relation_state()
    payload = {
        "source_offset": int(os.getenv("RELATION_SOURCE_OFFSET", "0")),
        "source_limit": int(os.getenv("RELATION_SOURCE_LIMIT", "0")),
        "source_cycle": int(os.getenv("RELATION_SOURCE_CYCLE", "0")),
        "people_observations": rows(workbook, "People"),
        "relationship_observations": rows(workbook, "Relationships"),
        "unique_entities": rows(workbook, "Unique_People"),
        "strict_scope_entities": rows(workbook, "GCC_Unique_People"),
        "all_source_qids": len(relation_state.get("all_source_qids", [])),
        "retained_relation_records": len(relation_state.get("records", [])),
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False))


if __name__ == "__main__":
    main()
