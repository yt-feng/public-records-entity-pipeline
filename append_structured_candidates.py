"""Append unseen public structured-family candidates to the observation layer."""

from __future__ import annotations

import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).parent
WORKBOOK = ROOT / "inputs/entity_network_master.xlsx"
INPUT = ROOT / "data/structured_candidate_records.json"
SOURCE_ID = "SRC-069"


def style(ws, source: int, target: int) -> None:
    for col in range(1, ws.max_column + 1):
        before, after = ws.cell(source, col), ws.cell(target, col)
        if before.has_style:
            after._style = copy(before._style)
        if before.alignment:
            after.alignment = copy(before.alignment)
        if before.protection:
            after.protection = copy(before.protection)


def main() -> None:
    records = json.loads(INPUT.read_text(encoding="utf-8"))["records"]
    wb = load_workbook(WORKBOOK)
    people, sources, attachment = wb["People"], wb["Sources"], wb["Attachment_Readout"]
    existing_ids = {str(row[0]).strip() for row in people.iter_rows(min_row=2, values_only=True) if row[0]}
    existing_urls = {str(row[14]).strip() for row in people.iter_rows(min_row=2, values_only=True) if len(row) > 14 and row[14]}
    adds = [r for r in records if r.get("record_id") not in existing_ids and r.get("page_url") not in existing_urls]
    row_number = people.max_row
    for record in adds:
        people.append([
            record["record_id"], record["country_section"], record["house"],
            f'{record["country_section"]}/structured family candidate', record["name"], "",
            "Public structured family candidate; current identity, role and status not assumed",
            record["intro"], record["relationship"], record["relation_type"], record["record_status"],
            "Public structured candidate layer", record["evidence_level"], SOURCE_ID, record["page_url"],
            record["as_of"], "Confirm identity, family membership and current status independently",
            f'Family QID={record["family_qid"]}', "",
        ])
        row_number += 1
        style(people, row_number - 1, row_number)
    people.tables["PeopleTable"].ref = f"A1:S{people.max_row}"

    source_ids = {row[0] for row in sources.iter_rows(min_row=2, values_only=True) if row[0]}
    if SOURCE_ID not in source_ids:
        sources.append([
            SOURCE_ID, "开放结构化数据", "Public structured family membership candidates",
            "https://query.wikidata.org/sparql",
            "读取公开 P53 家族/王朝成员字段；按 QID 去重并标为候选层",
            "低/结构化候选", "2026-09-20", "家族字段是候选线索，不等于当前身份、居住地或在世状态",
        ])
        style(sources, sources.max_row - 1, sources.max_row)
    sources.tables["SourcesTable"].ref = f"A1:H{sources.max_row}"

    attachment_ids = {row[0] for row in attachment.iter_rows(min_row=2, values_only=True) if row[0]}
    if SOURCE_ID not in attachment_ids:
        attachment.append([
            SOURCE_ID, "Public structured family membership candidates。", "候选扩展层",
            "公开 P53 家族/王朝成员字段；新增人物仅写入未覆盖的公开 Wikidata 页面 URL",
            "否", "结构化候选需独立核验",
        ])
        style(attachment, attachment.max_row - 1, attachment.max_row)
    attachment.tables["AttachmentReadoutTable"].ref = f"A1:F{attachment.max_row}"
    wb.save(WORKBOOK)
    print(json.dumps({"input_records": len(records), "new_people": len(adds), "people": people.max_row - 1}, ensure_ascii=False))


if __name__ == "__main__":
    main()
