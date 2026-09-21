"""Export durable, compressed public extracts beside the reader-facing XLSX.

GitHub rejects a single repository file larger than 100 MB.  The full workbook
remains a workflow artifact; these gzip CSV extracts keep the public records
and unique layers available in the repository without committing the oversized
binary workbook.
"""

from __future__ import annotations

import csv
import gzip
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).parent
WORKBOOK = ROOT / "inputs/entity_network_master.xlsx"
OUTPUTS = {
    "Unique_People": ROOT / "data/public_unique_people.csv.gz",
    "GCC_Unique_People": ROOT / "data/public_gcc_unique_people.csv.gz",
    "People": ROOT / "data/public_people_observations.csv.gz",
    "Relationships": ROOT / "data/public_relationships.csv.gz",
}
FORBIDDEN = re.compile(r"(?:/Users/|cloud-storage|public-researcher|workstream-0[12])", re.I)


def value_text(value: object) -> str:
    return "" if value is None else str(value)


def export_sheet(workbook, sheet_name: str, output: Path) -> tuple[int, int]:
    sheet = workbook[sheet_name]
    output.parent.mkdir(parents=True, exist_ok=True)
    temp = output.with_suffix(output.suffix + ".tmp")
    rows = 0
    columns = 0
    with gzip.open(temp, "wt", encoding="utf-8", newline="", compresslevel=6) as handle:
        writer = csv.writer(handle, lineterminator="\n")
        for row in sheet.iter_rows(values_only=True):
            values = [value_text(value) for value in row]
            if any(FORBIDDEN.search(value) for value in values):
                raise SystemExit(f"public export contains a private/local label in {sheet_name}")
            writer.writerow(values)
            rows += 1
            columns = max(columns, len(values))
    temp.replace(output)
    return rows, columns


def main() -> None:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    summary = {}
    try:
        for sheet_name, output in OUTPUTS.items():
            if sheet_name not in workbook.sheetnames:
                raise SystemExit(f"missing sheet for public export: {sheet_name}")
            rows, columns = export_sheet(workbook, sheet_name, output)
            summary[sheet_name] = {
                "rows": rows,
                "columns": columns,
                "bytes": output.stat().st_size,
                "output": str(output),
            }
    finally:
        workbook.close()
    print(summary)


if __name__ == "__main__":
    main()
