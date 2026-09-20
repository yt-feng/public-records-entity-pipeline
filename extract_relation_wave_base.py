"""Read one more explicit relationship wave from newly materialized GCC QIDs."""

from __future__ import annotations

import hashlib
import json
import random
import re
import time
import urllib.parse
import urllib.request
from pathlib import Path
from urllib.error import HTTPError

from openpyxl import load_workbook

ROOT = Path(__file__).parent
WORKBOOK = ROOT / "inputs/entity_network_master.xlsx"
PREVIOUS = ROOT / "data/wikidata_gcc_entity_claim_records.json"
OUTPUT = ROOT / "data/network_edge_wave2_records.json"
CHECKPOINT = ROOT / "data/wikidata_gcc_entity_claim_wave2_checkpoint.json"
SOURCE_ID = "SRC-062"
GCC = {"Saudi Arabia", "Qatar", "United Arab Emirates", "Kuwait", "Bahrain", "Oman"}
PROPERTIES = {"P22": "father", "P25": "mother", "P3373": "sibling", "P26": "spouse", "P40": "child"}


def batches(items: list[str], size: int = 50):
    for i in range(0, len(items), size): yield items[i:i + size]


def api(ids: list[str]) -> dict:
    query = urllib.parse.urlencode({"action": "wbgetentities", "ids": "|".join(ids), "props": "claims|labels", "languages": "en", "format": "json"})
    req = urllib.request.Request(f"https://www.wikidata.org/w/api.php?{query}", headers={"User-Agent": "Public-Records-Research/0.4"})
    for attempt in range(9):
        try:
            with urllib.request.urlopen(req, timeout=60) as response: return json.load(response)
        except HTTPError as exc:
            if exc.code != 429 or attempt == 8: raise
            retry_after = exc.headers.get("Retry-After") if exc.headers else None
            try:
                delay = float(retry_after) if retry_after else 15 * (attempt + 1)
            except (TypeError, ValueError):
                delay = 15 * (attempt + 1)
            # Spread concurrent retries so a burst does not immediately
            # recreate the same Wikidata rate-limit window.
            time.sleep(min(max(delay, 5.0) + random.uniform(0, 5), 180.0))
    raise RuntimeError("unreachable")


def label(entity: dict) -> str:
    return entity.get("labels", {}).get("en", {}).get("value", entity.get("id", ""))


def targets(entity: dict):
    for prop, edge in PROPERTIES.items():
        for claim in entity.get("claims", {}).get(prop, []):
            value = claim.get("mainsnak", {}).get("datavalue", {}).get("value", {})
            qid = value.get("id") if isinstance(value, dict) else None
            if qid: yield qid, edge


def stable_id(key: tuple[str, str, str, str, str]) -> str:
    return "WDC2-" + hashlib.sha1("|".join(key).encode("utf-8")).hexdigest()[:16].upper()


def main() -> None:
    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    contexts: dict[str, set[tuple[str, str]]] = {}
    for row in wb["People"].iter_rows(min_row=2, values_only=True):
        if row[1] not in GCC: continue
        m = re.search(r"wikidata\.org/(?:wiki/|entity/)(Q\d+)", str(row[14] or ""))
        if m: contexts.setdefault(m.group(1), set()).add((row[1], row[2]))
    source_qids = sorted(contexts)
    previous = json.loads(PREVIOUS.read_text(encoding="utf-8"))["records"] if PREVIOUS.exists() else []
    previous_keys = {(r["member_qid"], r["related_qid"], r["edge"], r["country_section"], r["house"]) for r in previous}
    if CHECKPOINT.exists():
        state = json.loads(CHECKPOINT.read_text(encoding="utf-8")); source_entities = state.get("source_entities", {}); done = set(state.get("done", []))
    else: source_entities, done = {}, set()
    for batch in batches([q for q in source_qids if q not in done]):
        source_entities.update(api(batch).get("entities", {})); done.update(batch)
        CHECKPOINT.write_text(json.dumps({"source_entities": source_entities, "done": sorted(done)}, ensure_ascii=False), encoding="utf-8"); time.sleep(0.2)
    target_qids = sorted({qid for e in source_entities.values() for qid, _ in targets(e)})
    target_entities: dict[str, dict] = {}
    target_checkpoint = ROOT / "data/wikidata_gcc_entity_claim_wave2_target_checkpoint.json"
    if target_checkpoint.exists():
        target_state=json.loads(target_checkpoint.read_text(encoding="utf-8")); target_entities=target_state.get("entities",{}); target_done=set(target_state.get("done",[]))
    else: target_done=set()
    for batch in batches([q for q in target_qids if q not in target_done]):
        target_entities.update(api(batch).get("entities", {})); target_done.update(batch)
        target_checkpoint.write_text(json.dumps({"entities":target_entities,"done":sorted(target_done)},ensure_ascii=False),encoding="utf-8"); time.sleep(0.2)
    records, seen = [], set()
    for member_qid, entity in source_entities.items():
        for related_qid, edge in targets(entity):
            target = target_entities.get(related_qid, {}); p31={c.get("mainsnak",{}).get("datavalue",{}).get("value",{}).get("id") for c in target.get("claims",{}).get("P31",[])}
            if "Q5" not in p31: continue
            for country, house in sorted(contexts.get(member_qid,set())):
                key=(member_qid,related_qid,edge,country,house)
                if key in previous_keys or key in seen: continue
                seen.add(key); member_name, related_name=label(entity),label(target)
                records.append({"record_id":stable_id(key),"country_section":country,"house":house,"member_qid":member_qid,"related_qid":related_qid,"member_name":member_name,"related_name":related_name,"edge":edge,"page_url":f"https://www.wikidata.org/wiki/{related_qid}","name":related_name,"source_text_short":f"Wikidata wave-2 {edge} claim from GCC entity {member_name}; endpoint is a Wikidata human entity.","record_status":"public GCC family-network candidate; second-wave explicit relationship claim; current status not assumed","evidence_level":"低/结构化关系扩展","intro":f"Wikidata wave 2 records {member_name} -> {related_name} as {edge}; endpoint P31 includes human; independent verification required.","relationship":f"{member_name} -> {related_name}: {edge}; source context={house}; second-wave QID expansion.","relation_type":f"GCC家族网络第二波/{edge}","source_id":SOURCE_ID,"as_of":"2026-09-19"})
    OUTPUT.write_text(json.dumps({"records":records,"source_qids":source_qids,"target_qids":target_qids},ensure_ascii=False,indent=2),encoding="utf-8")
    for p in [CHECKPOINT,target_checkpoint]:
        if p.exists(): p.unlink()
    print(json.dumps({"source_qids":len(source_qids),"target_qids":len(target_qids),"new_records":len(records),"output":str(OUTPUT)},ensure_ascii=False))


if __name__ == "__main__": main()
