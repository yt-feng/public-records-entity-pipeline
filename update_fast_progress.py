"""Update the public unique-entity progress index without rewriting the workbook.

The workbook is the reader-facing artifact.  Relation-wave JSON is the
append-only processing state.  Re-opening a 100k+ row XLSX for every small
relation wave is needlessly expensive, so this script initializes a compact
identity index from the workbook once and applies later observations to that
index.  The final workbook is materialized only after the target is reached.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from build_entity_network_layer import GCC, exclusion_reason
from relation_state import load_relation_state


ROOT = Path(__file__).parent
WORKBOOK = ROOT / "inputs/entity_network_master.xlsx"
INDEX = ROOT / "data/fast_entity_index.json"
PROGRESS = ROOT / "data/pipeline_progress.json"
TARGET_UNIQUE = 100_000


def split_values(value: object) -> list[str]:
    return [part.strip() for part in re.split(r"[；;]", str(value or "")) if part.strip()]


def entity_id_for_key(key: str) -> str:
    return "ENT-" + hashlib.sha1(key.encode("utf-8")).hexdigest()[:16].upper()


def entity_is_strict_gcc(entity: dict) -> bool:
    countries = set(entity.get("countries", []))
    return bool(countries) and countries <= GCC


def ensure_contexts(entity: dict) -> list[list[str]]:
    contexts = entity.setdefault("contexts", [])
    if contexts:
        return contexts
    for country in entity.get("countries", []):
        for house in entity.get("houses", []) or [""]:
            pair = [country, house]
            if pair not in contexts:
                contexts.append(pair)
    return contexts


def initialize_index() -> dict:
    """Create the initial index from the latest committed public workbook."""
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    entities: dict[str, dict] = {}
    key_to_entity: dict[str, str] = {}
    for row in workbook["Unique_People"].iter_rows(min_row=2, values_only=True):
        entity_id = str(row[0] or "").strip()
        if not entity_id:
            continue
        entity = {
            "keys": split_values(row[6]),
            "countries": split_values(row[3]),
            "houses": split_values(row[4]),
            "names": split_values(row[2] or row[1]),
            "contexts": [],
        }
        ensure_contexts(entity)
        entities[entity_id] = entity
        for key in entity["keys"]:
            key_to_entity[key] = entity_id

    excluded_qids: set[str] = set()
    if "Excluded_Candidates" in workbook.sheetnames:
        for row in workbook["Excluded_Candidates"].iter_rows(min_row=2, values_only=True):
            match = re.search(r"wikidata\.org/(?:wiki/|entity/)(Q\d+)", str(row[8] or ""))
            if match:
                excluded_qids.add(match.group(1))

    observation_ids = {
        str(row[0]).strip()
        for row in workbook["People"].iter_rows(min_row=2, max_col=1, values_only=True)
        if row[0]
    }
    people_observations = max(workbook["People"].max_row - 1, 0)
    relationship_observations = max(workbook["Relationships"].max_row - 1, 0)
    workbook.close()
    return {
        "version": 1,
        "entities": entities,
        "key_to_entity": key_to_entity,
        "excluded_qids": sorted(excluded_qids),
        "observation_ids": sorted(observation_ids),
        "people_observations": people_observations,
        "relationship_observations": relationship_observations,
    }


def load_or_initialize() -> dict:
    if INDEX.exists():
        return json.loads(INDEX.read_text(encoding="utf-8"))
    return initialize_index()


def main() -> None:
    state = load_or_initialize()
    entities: dict[str, dict] = state["entities"]
    key_to_entity: dict[str, str] = state["key_to_entity"]
    excluded_qids = set(state.get("excluded_qids", []))
    observation_ids = set(state.get("observation_ids", []))
    people_observations = int(state.get("people_observations", 0))
    relationship_observations = int(state.get("relationship_observations", 0))

    relation_state = load_relation_state()
    new_observations = 0
    new_entities = 0
    for record in relation_state.get("records", []):
        record_id = str(record.get("record_id") or "").strip()
        if not record_id or record_id in observation_ids:
            continue
        observation_ids.add(record_id)
        people_observations += 1
        relationship_observations += 1
        new_observations += 1

        related_qid = str(record.get("related_qid") or "").strip()
        name = str(record.get("name") or record.get("related_name") or "").strip()
        page_url = str(record.get("page_url") or "").strip()
        country = str(record.get("country_section") or "").strip()
        house = str(record.get("house") or "").strip()
        if not related_qid or related_qid in excluded_qids:
            continue
        if exclusion_reason([name], [page_url], [country, house]):
            continue

        key = f"QID:{related_qid}"
        entity_id = key_to_entity.get(key)
        if entity_id is None:
            entity_id = entity_id_for_key(key)
            entities[entity_id] = {"keys": [key], "countries": [], "houses": [], "names": [], "contexts": []}
            key_to_entity[key] = entity_id
            new_entities += 1
        entity = entities[entity_id]
        ensure_contexts(entity)
        for field, value in (("countries", country), ("houses", house), ("names", name)):
            if value and value not in entity[field]:
                entity[field].append(value)
        context = [country, house]
        if context[0] and context not in entity["contexts"]:
            entity["contexts"].append(context)

    unique_entities = len(entities)
    strict_scope_entities = sum(1 for entity in entities.values() if entity_is_strict_gcc(entity))
    state.update(
        {
            "version": 1,
            "entities": entities,
            "key_to_entity": key_to_entity,
            "excluded_qids": sorted(excluded_qids),
            "observation_ids": sorted(observation_ids),
            "people_observations": people_observations,
            "relationship_observations": relationship_observations,
        }
    )
    INDEX.write_text(json.dumps(state, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")

    progress = {
        "source_offset": int(os.getenv("RELATION_SOURCE_OFFSET", "0")),
        "source_limit": int(os.getenv("RELATION_SOURCE_LIMIT", "0")),
        "source_cycle": int(os.getenv("RELATION_SOURCE_CYCLE", "0")),
        "people_observations": people_observations,
        "relationship_observations": relationship_observations,
        "unique_entities": unique_entities,
        "strict_scope_entities": strict_scope_entities,
        "all_source_qids": len(relation_state.get("all_source_qids", [])),
        "retained_relation_records": len(relation_state.get("records", [])),
        "unique_gap_to_100k": max(TARGET_UNIQUE - unique_entities, 0),
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }
    PROGRESS.write_text(json.dumps(progress, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        json.dumps(
            {
                **progress,
                "new_observations": new_observations,
                "new_entities": new_entities,
                "index": str(INDEX),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
