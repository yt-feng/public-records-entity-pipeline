"""Create a filtered GCC-only view from the conservative Unique_People layer."""

from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).parent
WORKBOOK = ROOT / "inputs/entity_network_master.xlsx"
GCC = {"Saudi Arabia", "Qatar", "United Arab Emirates", "Kuwait", "Bahrain", "Oman"}


def main() -> None:
    wb = load_workbook(WORKBOOK)
    source = wb["Unique_People"]
    if "GCC_Unique_People" in wb.sheetnames:
        del wb["GCC_Unique_People"]
    ws = wb.create_sheet("GCC_Unique_People", 2)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    headers = [cell.value for cell in source[1]]
    ws.append(headers)
    rows = []
    for row in source.iter_rows(min_row=2, values_only=True):
        countries = {value.strip() for value in str(row[3] or "").split("；") if value.strip()}
        # Strict GCC view: a person is included only when every country label
        # attached to the conservative entity is one of the six GCC members.
        # This prevents an Oman relationship edge with Afghanistan, Iran,
        # India/South Asia, or another historical scope from being presented as
        # a GCC person merely because one value happens to be Oman.
        if countries and countries <= GCC:
            rows.append(row)
    for row in rows:
        ws.append(list(row))
    for col in range(1, source.max_column + 1):
        template = source.cell(1, col)
        target = ws.cell(1, col)
        target.font = copy(template.font)
        target.fill = copy(template.fill)
        target.border = copy(template.border)
        target.alignment = copy(template.alignment)
        target.number_format = template.number_format
        ws.column_dimensions[target.column_letter].width = source.column_dimensions[target.column_letter].width
    ws.row_dimensions[1].height = source.row_dimensions[1].height
    table = Table(displayName="GCCUniquePeopleTable", ref=f"A1:T{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)
    wb["ReadMe"]["B6"] = f"GCC_Unique_People 为严格 GCC 六国视图：只有国家字段全部属于 Saudi Arabia、Qatar、United Arab Emirates、Kuwait、Bahrain、Oman 才纳入，共 {len(rows):,} 个保守去重实体；跨入非 GCC 历史/亲缘范围的记录留在 Unique_People，但不计入 GCC 视图。"
    wb.save(WORKBOOK)
    print({"gcc_unique_people": len(rows), "sheet": "GCC_Unique_People", "output": str(WORKBOOK)})


if __name__ == "__main__":
    main()
